"""文档解析：把各类文件提取为带页码的文本段落。

统一输出格式：list[段文本]，每段附带来源页码（纯文本文档页码为 0）。
"""

import io
from dataclasses import dataclass

from pypdf import PdfReader


@dataclass
class ParsedSegment:
    text: str
    page: int


def _pdf(data: bytes) -> list[ParsedSegment]:
    reader = PdfReader(io.BytesIO(data))
    segments: list[ParsedSegment] = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            segments.append(ParsedSegment(text=text, page=i))
    return segments


def _docx(data: bytes) -> list[ParsedSegment]:
    import docx

    document = docx.Document(io.BytesIO(data))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    if not paragraphs:
        return []
    # docx 无固定分页，按段落顺序合并为单一分段，页码 0
    return [ParsedSegment(text="\n".join(paragraphs), page=0)]


def _xlsx(data: bytes) -> list[ParsedSegment]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    segments: list[ParsedSegment] = []
    for sheet in wb.worksheets:
        rows: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            segments.append(ParsedSegment(text=f"[工作表 {sheet.title}]\n" + "\n".join(rows), page=0))
    return segments


def _plain(data: bytes) -> list[ParsedSegment]:
    text = data.decode("utf-8", errors="replace").strip()
    return [ParsedSegment(text=text, page=0)] if text else []


PARSERS = {
    "pdf": _pdf,
    "docx": _docx,
    "xlsx": _xlsx,
    "md": _plain,
    "txt": _plain,
}


def parse_file(file_type: str, data: bytes) -> list[ParsedSegment]:
    parser = PARSERS.get(file_type)
    if parser is None:
        raise ValueError(f"不支持的文件类型: {file_type}")
    return parser(data)
